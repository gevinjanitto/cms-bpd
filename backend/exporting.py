"""Excel-friendly, localized reporting with matching CSV and XLSX columns."""
import csv
import io
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

STATUS = {'passed': 'Lulus', 'failed': 'Belum Lulus', 'reviewing': 'Menunggu Penilaian'}
WITA = ZoneInfo('Asia/Makassar')

def safe(value):
    if not isinstance(value, str): return value
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else value

def local_date(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(WITA).replace(tzinfo=None) if value else None

def report_data(users, quizzes, results, pairs, pending, view, search):
    if view in ['units', 'top', 'bottom', 'never']:
        rows = []
        if view == 'units':
            headers = ['No.', 'Unit Kerja', 'Karyawan', 'Penugasan', 'Selesai', 'Belum Mengikuti', 'Partisipasi (%)', 'Rata-rata Nilai']
            for unit in sorted({u['unit'] for u in users}):
                members = {u['id'] for u in users if u['unit'] == unit}
                total = sum(p[1] in members for p in pairs)
                done = [r for r in results if r['user_id'] in members]
                graded = [r for r in done if r['status'] != 'reviewing']
                rows.append([unit, len(members), total, len(done), total-len(done), round(len(done)/total*100, 1) if total else 0, round(sum(r['score'] for r in graded)/len(graded), 1) if graded else 0])
        else:
            headers = ['No.', 'Nama Karyawan', 'Unit Kerja', 'Jabatan', 'Penugasan', 'Selesai', 'Rata-rata Nilai']
            for user in users:
                done = [r for r in results if r['user_id'] == user['id']]
                total = sum(p[1] == user['id'] for p in pairs)
                graded = [r for r in done if r['status'] != 'reviewing']
                avg = round(sum(r['score'] for r in graded)/len(graded), 1) if graded else 0
                if (view == 'never' and not done and total) or (view != 'never' and done):
                    rows.append([user['name'], user['unit'], user['position'], total, len(done), avg])
            if view != 'never': rows.sort(key=lambda r: r[-1], reverse=view == 'top')
            rows = rows[:10]
    else:
        headers = ['No.', 'Nama Karyawan', 'Unit Kerja', 'Kuis', 'Nilai', 'Passing Grade', 'Status', 'Waktu Pengumpulan (WITA)']
        rows = []
        if view != 'pending':
            for r in results:
                if view == 'essay' and r['status'] != 'reviewing': continue
                rows.append([r['user_name'], r['unit'], r['quiz_title'], r['score'], r['passing_grade'], STATUS.get(r['status'], r['status']), local_date(r['submitted_at'])])
        if view in ['all', 'pending']:
            for r in pending:
                rows.append([r['user_name'], r['unit'], r['quiz_title'], None, None, 'Belum Mengikuti', None])
    rows = [r for r in rows if not search or search.lower() in ' '.join(str(v) for v in r).lower()]
    return headers, [[i + 1, *[safe(v) for v in row]] for i, row in enumerate(rows)]


def csv_bytes(headers, rows):
    output = io.StringIO(newline='')
    output.write('sep=;\r\n')
    writer = csv.writer(output, delimiter=';', lineterminator='\r\n', quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([v.strftime('%d/%m/%Y %H:%M:%S') if isinstance(v, datetime) else str(v).replace('.', ',') if isinstance(v, float) else v for v in row])
    return ('\ufeff' + output.getvalue()).encode('utf-8')


def xlsx_bytes(headers, rows, period, unit):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Laporan Kepatuhan'
    last = get_column_letter(len(headers))
    for row, text in [(1, 'COMPLIANCE MANAGEMENT SYSTEM'), (2, 'Bank BPD Bali · Laporan Kepatuhan'), (3, f"Periode: {period or 'Semua'}  |  Unit: {unit or 'Semua Unit Kerja'}  |  Waktu: WITA")]:
        sheet.merge_cells(f'A{row}:{last}{row}')
        sheet.cell(row, 1, text).font = Font(name='Calibri', size=16 if row == 1 else 11, bold=row <= 2, color='187C57' if row == 1 else '44594B')
        sheet.row_dimensions[row].height = 27 if row == 1 else 23
    for c, title in enumerate(headers, 1):
        cell = sheet.cell(5, c, title)
        cell.fill = PatternFill('solid', fgColor='187C57')
        cell.font = Font(name='Calibri', size=11, color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[5].height = 32
    for r, values in enumerate(rows, 6):
        sheet.row_dimensions[r].height = 32
        for c, value in enumerate(values, 1):
            cell = sheet.cell(r, c, value)
            cell.font = Font(name='Calibri', size=11, color='26382F')
            cell.fill = PatternFill('solid', fgColor='F0F7F2' if r % 2 == 0 else 'FFFFFF')
            cell.border = Border(bottom=Side(style='hair', color='DCE7DF'))
            cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center' if c == 1 or isinstance(value, (int, float)) else 'left')
            if isinstance(value, datetime): cell.number_format = 'dd/mm/yyyy hh:mm:ss'
            elif isinstance(value, (int, float)): cell.number_format = '0.#'
    widths = {'No.': 6, 'Nama Karyawan': 30, 'Unit Kerja': 30, 'Kuis': 44, 'Status': 24, 'Waktu Pengumpulan (WITA)': 28, 'Jabatan': 22}
    for c, header in enumerate(headers, 1): sheet.column_dimensions[get_column_letter(c)].width = widths.get(header, max(15, len(header) + 3))
    sheet.freeze_panes = 'A6'
    sheet.auto_filter.ref = f'A5:{last}{max(5, len(rows) + 5)}'
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:5'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()